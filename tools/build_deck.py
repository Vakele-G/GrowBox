#!/usr/bin/env python3
"""Generate the GrowBox hackathon submission deck.

Creates a 16:9 PowerPoint deck styled with the GrowBox brand palette
(ColorHunt E3F2FD / 90CAF9 / 2196F3 / 0D47A1, matching the live app):

    python3 tools/build_deck.py [output.pptx]

Content sources: README.md, PROJECT_SPEC.md, SUBMISSION.md, MILESTONES.md.
The deck is a generated artifact; re-run this script after editing it.
"""

from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
LOGO = ROOT / "public" / "icons" / "icon-512.png"
DEFAULT_OUT = ROOT / "GrowBox_Submission_Deck.pptx"

# ---------------------------------------------------------------- palette ---
BRAND = RGBColor(0x0D, 0x47, 0xA1)  # primary: headings, dark fills
DEEP = RGBColor(0x0A, 0x36, 0x79)  # gradient start
BLUE = RGBColor(0x21, 0x96, 0xF3)  # accent
SKY = RGBColor(0x90, 0xCA, 0xF9)  # soft chips, tints
MIST = RGBColor(0xE3, 0xF2, 0xFD)  # panel background
LINE = RGBColor(0xC3, 0xDD, 0xF5)  # borders (derived from sky)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x10, 0x22, 0x3D)  # body text
MUTED = RGBColor(0x4A, 0x60, 0x80)  # secondary text

FONT = "Segoe UI"
SLIDE_W = 13.333
SLIDE_H = 7.5
FOOTER = "GrowBox · FNB App of the Year 2026 · grow-box.vercel.app"


# ------------------------------------------------------------- primitives ---
def blank(prs: Presentation):
    return prs.slides.add_slide(prs.slide_layouts[6])


def rect(slide, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, radius=None):
    sh = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(1)
    sh.shadow.inherit = False
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            sh.adjustments[0] = radius
        except (IndexError, ValueError):
            pass
    return sh


def gradient_bg(slide, angle: float = 45.0):
    bg = rect(slide, 0, 0, SLIDE_W, SLIDE_H, fill=DEEP)
    bg.fill.gradient()
    bg.fill.gradient_angle = angle
    stops = bg.fill.gradient_stops
    stops[0].color.rgb = DEEP
    stops[0].position = 0.0
    stops[1].color.rgb = BLUE
    stops[1].position = 1.0
    return bg


def textbox(slide, x, y, w, h, name=None, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    if name:
        tb.name = name
    return tf


def para(
    tf,
    text,
    *,
    size,
    color,
    bold=False,
    first=False,
    align=PP_ALIGN.LEFT,
    bullet=False,
    bullet_color=None,
    space_before=0,
    space_after=6,
    line_spacing=1.05,
):
    """Add a paragraph. `text` is a str or list of (str, color, bold) runs."""
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.line_spacing = line_spacing
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)

    runs = [(text, color, bold)] if isinstance(text, str) else text
    for value, run_color, run_bold in runs:
        run = p.add_run()
        run.text = value
        run.font.size = Pt(size)
        run.font.name = FONT
        run.font.bold = run_bold
        run.font.color.rgb = run_color

    if bullet:
        # schema order inside a:pPr: spacing -> buClr -> buFont -> buChar
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(Inches(0.24)))
        pPr.set("indent", str(-Inches(0.24)))
        bullet_hex = bullet_color or BLUE
        pPr.append(
            parse_xml(
                f'<a:buClr {nsdecls("a")}><a:srgbClr val="{bullet_hex}"/></a:buClr>'
            )
        )
        pPr.append(parse_xml(f'<a:buFont {nsdecls("a")} typeface="Arial"/>'))
        pPr.append(parse_xml(f'<a:buChar {nsdecls("a")} char="\u25b8"/>'))
    return p


def bullets(slide, x, y, w, h, items, *, size=13.5, color=INK, gap=9, name=None):
    tf = textbox(slide, x, y, w, h, name=name)
    for i, item in enumerate(items):
        para(
            tf,
            item,
            size=size,
            color=color,
            first=(i == 0),
            bullet=True,
            bullet_color=BLUE,
            space_after=gap,
            line_spacing=1.08,
        )
    return tf


def stat_card(slide, x, y, w, h, big, label, *, fill=MIST, big_color=BRAND):
    rect(slide, x, y, w, h, fill=fill, line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
    tf = textbox(slide, x + 0.12, y + 0.1, w - 0.24, h - 0.2, anchor=MSO_ANCHOR.MIDDLE)
    para(tf, big, size=30, color=big_color, bold=True, first=True, align=PP_ALIGN.CENTER, space_after=2)
    para(tf, label, size=10.5, color=MUTED, align=PP_ALIGN.CENTER, space_after=0, line_spacing=1.0)


def card(slide, x, y, w, h, title, body, *, accent=BRAND, number=None, body_size=10.5):
    rect(slide, x, y, w, h, fill=WHITE, line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.07)
    top = y + 0.18
    if number is not None:
        chip = rect(slide, x + 0.18, top, 0.4, 0.4, fill=accent, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.25)
        ctf = chip.text_frame
        ctf.word_wrap = False
        ctf.margin_left = ctf.margin_right = ctf.margin_top = ctf.margin_bottom = 0
        para(ctf, str(number), size=13, color=WHITE, bold=True, first=True, align=PP_ALIGN.CENTER, space_after=0)
        ctf.vertical_anchor = MSO_ANCHOR.MIDDLE
        top += 0.56
    tf = textbox(slide, x + 0.18, top, w - 0.36, y + h - top - 0.14)
    para(tf, title, size=13, color=accent, bold=True, first=True, space_after=5, line_spacing=1.0)
    para(tf, body, size=body_size, color=MUTED, space_after=0, line_spacing=1.1)


def band(slide, x, y, w, h, text, *, fill=BRAND, color=WHITE, size=15):
    rect(slide, x, y, w, h, fill=fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
    tf = textbox(slide, x + 0.2, y, w - 0.4, h, anchor=MSO_ANCHOR.MIDDLE)
    para(tf, text, size=size, color=color, bold=True, first=True, align=PP_ALIGN.CENTER, space_after=0)


def chip(slide, x, y, w, h, text, *, fill=MIST, color=BRAND, size=13):
    rect(slide, x, y, w, h, fill=fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
    tf = textbox(slide, x + 0.1, y, w - 0.2, h, anchor=MSO_ANCHOR.MIDDLE)
    para(tf, text, size=size, color=color, bold=True, first=True, align=PP_ALIGN.CENTER, space_after=0)


# ------------------------------------------------------------ slide chrome ---
def content_slide(prs, kicker: str, title: str, number: int):
    slide = blank(prs)
    rect(slide, 0, 0, 0.16, SLIDE_H, fill=BRAND)
    tf = textbox(slide, 0.8, 0.46, 11.9, 0.3)
    para(tf, kicker.upper(), size=11, color=BLUE, bold=True, first=True, space_after=0)
    tf = textbox(slide, 0.8, 0.76, 11.9, 0.7)
    para(tf, title, size=30, color=BRAND, bold=True, first=True, space_after=0)
    rect(slide, 0.82, 1.44, 1.5, 0.055, fill=BLUE)
    tf = textbox(slide, 0.8, 7.04, 9.0, 0.3)
    para(tf, FOOTER, size=9.5, color=MUTED, first=True, space_after=0)
    tf = textbox(slide, 11.9, 7.04, 0.7, 0.3)
    para(tf, str(number), size=9.5, color=MUTED, first=True, align=PP_ALIGN.RIGHT, space_after=0)
    return slide


def hero_slide(prs):
    """Shared full-bleed background for the title and closing slides."""
    slide = blank(prs)
    gradient_bg(slide, angle=30.0)
    return slide


def logo(slide, x, y, size, *, card_size=None):
    if card_size:
        rect(slide, x, y, card_size, card_size, fill=WHITE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.16)
        pad = (card_size - size) / 2
        slide.shapes.add_picture(str(LOGO), Inches(x + pad), Inches(y + pad), width=Inches(size))
    else:
        slide.shapes.add_picture(str(LOGO), Inches(x), Inches(y), width=Inches(size))


# ------------------------------------------------------------------ slides ---
def build(out_path: Path) -> Presentation:
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)

    # -- 1 · Title ----------------------------------------------------------
    s = hero_slide(prs)
    logo(s, 5.82, 1.15, 1.3, card_size=1.7)
    tf = textbox(s, 1.5, 3.05, 10.33, 1.3)
    para(tf, "GrowBox", size=64, color=WHITE, bold=True, first=True, align=PP_ALIGN.CENTER, space_after=0)
    tf = textbox(s, 1.5, 4.3, 10.33, 0.62)
    para(
        tf,
        "From a street-side hustle to a bankable business.",
        size=22,
        color=MIST,
        first=True,
        align=PP_ALIGN.CENTER,
        space_after=0,
    )
    tf = textbox(s, 1.9, 4.95, 9.53, 0.45)
    para(
        tf,
        "A proactive business mentor for South African informal entrepreneurs.",
        size=14,
        color=SKY,
        first=True,
        align=PP_ALIGN.CENTER,
        space_after=0,
    )
    chip(s, 4.67, 5.6, 4.0, 0.52, "grow-box.vercel.app", fill=MIST, color=BRAND, size=15)
    tf = textbox(s, 1.5, 6.7, 10.33, 0.4)
    para(
        tf,
        "FNB App of the Year Hackathon 2026 · Youth entrepreneurship and small-business growth",
        size=11.5,
        color=SKY,
        first=True,
        align=PP_ALIGN.CENTER,
        space_after=0,
    )

    # -- 2 · Problem --------------------------------------------------------
    s = content_slide(prs, "Problem", "The concrete problem we solve", 2)
    stats = [
        ("1.9M", "non-VAT-registered\nbusinesses"),
        ("80%", "are unregistered"),
        ("77%", "manage money by hand"),
        ("33%", "can access credit"),
    ]
    x = 0.82
    for big, label in stats:
        stat_card(s, x, 1.8, 2.82, 1.5, big, label)
        x += 2.98
    bullets(
        s,
        0.82,
        3.6,
        7.4,
        2.2,
        [
            "No paper trail — so no credit.",
            "Margins collapse when input prices move.",
            "Advice is generic, scattered and reactive.",
            "R350bn MSME funding gap (OECD, 2026).",
        ],
        size=15,
        gap=11,
    )
    band(s, 8.5, 3.6, 4.0, 2.2, "No paper trail\n= no credit.", fill=BRAND, size=22)

    # -- 3 · Solution -------------------------------------------------------
    s = content_slide(prs, "Solution", "Our unique solution", 3)
    band(
        s,
        0.82,
        1.8,
        11.7,
        0.85,
        "\u201cWhat is the next best step I should take?\u201d",
        fill=MIST,
        color=BRAND,
        size=20,
    )
    bullets(
        s,
        0.82,
        3.0,
        7.2,
        3.3,
        [
            "A proactive mentor — she never asks.",
            "Acts on her own numbers, every day.",
            "Answers in English and isiZulu.",
            "Builds the bankable record as a by-product.",
            "Live at grow-box.vercel.app.",
        ],
        size=15,
        gap=13,
    )
    rect(s, 8.5, 3.0, 4.0, 3.3, fill=MIST, line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
    tf = textbox(s, 8.8, 3.28, 3.4, 2.8)
    para(tf, "WHAT WE ARE NOT", size=11.5, color=BLUE, bold=True, first=True, space_after=12)
    for line in ("Not a chatbot.", "Not accounting software.", "Not another course."):
        para(tf, line, size=15, color=BRAND, bold=True, space_after=12, line_spacing=1.0)

    # -- 4 · Architecture ---------------------------------------------------
    s = content_slide(prs, "Architecture", "Our unique solution architecture", 4)
    layers = [
        ("PWA UI", "Next.js App Router · screens · server actions", MIST, INK),
        ("Pure-TS domain core", "costing · margins · runway · stage · score", SKY, INK),
        ("AI adapter (Gemini)", "explains and translates · 8s timeout · fallback", BLUE, WHITE),
        ("Supabase Postgres", "row-level security · money as integer cents", BRAND, WHITE),
        ("Vercel", "edge deploy · public production URL", DEEP, WHITE),
    ]
    y = 1.78
    for i, (name, detail, fill, text_color) in enumerate(layers):
        rect(s, 0.82, y, 7.4, 0.86, fill=fill, line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
        tf = textbox(s, 1.1, y + 0.12, 6.9, 0.62, anchor=MSO_ANCHOR.MIDDLE)
        para(tf, [(name + "  —  ", text_color, True), (detail, text_color, False)], size=13.5, color=text_color, first=True, space_after=0)
        if i < len(layers) - 1:
            arrow = rect(s, 4.32, y + 0.87, 0.4, 0.2, fill=SKY, shape=MSO_SHAPE.DOWN_ARROW)
            arrow.line.fill.background()
        y += 1.07
    notes = [
        ("Code calculates, AI explains.", "Every number is deterministic, tested TypeScript."),
        ("AI never blocks a flow.", "Rule-based fallback if the key or network is missing."),
        ("Rows are locked to one owner.", "RLS scopes every table to the business."),
    ]
    ny = 1.78
    for title, body in notes:
        card(s, 8.6, ny, 3.9, 1.5, title, body, accent=BRAND, body_size=11)
        ny += 1.7

    # -- 5 · Journey --------------------------------------------------------
    s = content_slide(prs, "Journey", "User journey story", 5)
    stages = ["1 · Idea", "2 · Launch", "3 · Stable", "4 · Grow", "5 · Enterprise"]
    x = 0.82
    for label in stages:
        chev = rect(s, x, 1.78, 2.36, 0.92, fill=BRAND, shape=MSO_SHAPE.CHEVRON)
        chev.line.fill.background()
        tf = chev.text_frame
        tf.word_wrap = False
        tf.margin_left = tf.margin_right = Inches(0.1)
        para(tf, label, size=14, color=WHITE, bold=True, first=True, align=PP_ALIGN.CENTER, space_after=0)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        x += 2.38
    bullets(
        s,
        0.82,
        3.2,
        11.7,
        3.2,
        [
            "Nomsa starts with R1,500 and an idea.",
            "The Stage 1 wizard costs her amagwinya in 90 seconds.",
            "She types: \u201csold 40 gwinyas, bought oil for R180\u201d.",
            "Nudges protect margin and cash before Thursday.",
            "14 days of trades become a lender-ready record.",
        ],
        size=16,
        gap=14,
    )

    # -- 6 · Product --------------------------------------------------------
    s = content_slide(prs, "Product", "What the app does", 6)
    features = [
        ("Costing wizard", "Cost per unit, margin, break-even, suggested price."),
        ("15-second logging", "Type a sale. Confirm the parse. Saved."),
        ("Proactive nudges", "Margin drops, cash warnings, stock cover."),
        ("Health Score", "A 0–100 credit-readiness signal."),
        ("Credit record", "Exportable trading history for lenders."),
    ]
    x = 0.82
    for i, (title, body) in enumerate(features, start=1):
        card(s, x, 1.9, 2.25, 2.7, title, body, accent=BLUE if i % 2 else BRAND, number=i)
        x += 2.38
    band(
        s,
        0.82,
        5.1,
        11.7,
        1.1,
        "Simulated FNB feed is labelled · confirm before anything is saved",
        fill=MIST,
        color=BRAND,
        size=15,
    )

    # -- 7 · Tech stack -----------------------------------------------------
    s = content_slide(prs, "Technology", "Key tech stack — and what makes it fast", 7)
    rect(s, 0.82, 1.85, 5.7, 4.5, fill=MIST, line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
    tf = textbox(s, 1.12, 2.1, 5.1, 0.4)
    para(tf, "STACK", size=12, color=BLUE, bold=True, first=True, space_after=0)
    bullets(
        s,
        1.12,
        2.6,
        5.1,
        3.6,
        [
            "Next.js 16 · App Router · Turbopack",
            "TypeScript strict · Tailwind CSS 4",
            "Supabase Postgres with row-level security",
            "Vercel edge hosting",
            "Vitest: 62 passing tests · ESLint · tsc",
        ],
        size=13.5,
        gap=12,
    )
    rect(s, 6.8, 1.85, 5.7, 4.5, fill=WHITE, line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
    tf = textbox(s, 7.1, 2.1, 5.1, 0.4)
    para(tf, "FAST BY DESIGN", size=12, color=BLUE, bold=True, first=True, space_after=0)
    bullets(
        s,
        7.1,
        2.6,
        5.1,
        3.6,
        [
            "Static-first shell, served from the edge",
            "Deterministic maths stays off the AI path",
            "Low-data PWA for entry-level Android",
            "Works offline for review and calculations",
            "English and isiZulu message files",
        ],
        size=13.5,
        gap=12,
    )

    # -- 8 · Demo account ---------------------------------------------------
    s = content_slide(prs, "Demo", "Demo account details", 8)
    rect(s, 0.82, 1.8, 6.0, 4.5, fill=MIST, line=BRAND, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
    tf = textbox(s, 1.12, 2.08, 5.4, 4.0)
    para(tf, "FOR JUDGES", size=11.5, color=BLUE, bold=True, first=True, space_after=12)
    para(tf, "Link", size=11, color=MUTED, bold=True, space_after=2)
    para(tf, "https://grow-box.vercel.app/login", size=16, color=BRAND, bold=True, space_after=14, line_spacing=1.0)
    para(tf, "One click", size=11, color=MUTED, bold=True, space_after=2)
    para(tf, "\u201cContinue with the demo account\u201d", size=16, color=BRAND, bold=True, space_after=14, line_spacing=1.0)
    para(tf, "Email (magic link, no password)", size=11, color=MUTED, bold=True, space_after=2)
    para(tf, "demo@growbox.app", size=16, color=BRAND, bold=True, space_after=0, line_spacing=1.0)
    tf = textbox(s, 7.2, 1.8, 5.3, 0.4)
    para(tf, "WHAT JUDGES WILL SEE", size=11.5, color=BLUE, bold=True, first=True, space_after=0)
    bullets(
        s,
        7.2,
        2.35,
        5.3,
        4.0,
        [
            "Nomsa's Kitchen — Soweto food business",
            "Stage 1 launch and costing wizard first",
            "Persona switcher loads 30 days of Stage 3 trades",
            "Logging, nudges, dashboard, Health Score",
            "English and isiZulu throughout",
        ],
        size=14,
        gap=14,
    )

    # -- 9 · App build file -------------------------------------------------
    s = content_slide(prs, "Distribution", "App build file — installable PWA", 9)
    bullets(
        s,
        0.82,
        1.9,
        7.0,
        4.3,
        [
            "One link installs it — no app store.",
            "Chrome on Android: Add to Home screen.",
            "Standalone window, own icon, portrait.",
            "Service worker caches pages offline.",
            "Updates instantly — no APK versions.",
        ],
        size=16,
        gap=18,
    )
    rect(s, 8.3, 1.9, 4.2, 4.3, fill=MIST, line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
    tf = textbox(s, 8.6, 2.16, 3.6, 3.8)
    para(tf, "BUILT FOR THE USER", size=11.5, color=BLUE, bold=True, first=True, space_after=14)
    for line, detail in (
        ("R399 Android", "entry-level hardware"),
        ("Prepaid data", "R30–R100 a month"),
        ("WhatsApp-first", "share a link, not a store page"),
    ):
        para(tf, line, size=16, color=BRAND, bold=True, space_after=2, line_spacing=1.0)
        para(tf, detail, size=12, color=MUTED, space_after=14, line_spacing=1.0)
    chip(s, 8.6, 5.5, 3.6, 0.5, "grow-box.vercel.app", fill=WHITE, color=BRAND, size=14)

    # -- 10 · Business model ------------------------------------------------
    s = content_slide(prs, "Business model", "How GrowBox makes money", 10)
    streams = [
        ("Free tier — R0", "Logging, dashboard, 3 nudges a week"),
        ("GrowBox Pro — R39/mo", "Full nudges, forecasting, exports"),
        ("Lender packs — R120", "Consented, verified trading history"),
        ("FNB referral — R250", "Per account or credit conversion"),
        ("B2B2C licence — R45k/mo", "DFIs, NYDA, TVET colleges, NGOs"),
        ("Marketplace — 3%", "Wholesale ingredient orders"),
    ]
    for i, (title, body) in enumerate(streams):
        col = i % 3
        row = i // 3
        card(s, 0.82 + col * 3.98, 1.85 + row * 1.7, 3.8, 1.5, title, body, accent=BRAND if row == 0 else BLUE, body_size=11)
    band(
        s,
        0.82,
        5.5,
        11.7,
        1.0,
        "Revenue: R1.08m (Y1)  →  R8.7m (Y2)  →  R31.6m (Y3)",
        fill=BRAND,
        size=18,
    )

    # -- 11 · Market & metrics ---------------------------------------------
    s = content_slide(prs, "Market", "Market and the metrics we track", 11)
    band(
        s,
        0.82,
        1.8,
        11.7,
        0.85,
        "North Star: Weekly Active Trading Businesses",
        fill=MIST,
        color=BRAND,
        size=18,
    )
    metrics = [
        ("25,000", "businesses in 12 months"),
        ("5,000", "bankable scores in 12 months"),
        ("150", "FNB conversions in year 1"),
        ("60%", "nudges acted on in 7 days"),
    ]
    x = 0.82
    for big, label in metrics:
        stat_card(s, x, 3.0, 2.82, 1.5, big, label)
        x += 2.98
    bullets(
        s,
        0.82,
        4.9,
        11.7,
        1.8,
        [
            "Launch where informal trade concentrates: Gauteng, KZN, Limpopo.",
            "WhatsApp-first sharing — near-zero acquisition cost.",
        ],
        size=15,
        gap=12,
    )

    # -- 12 · Why FNB -------------------------------------------------------
    s = content_slide(prs, "Partnership", "Why this matters to FNB", 12)
    bullets(
        s,
        0.82,
        1.95,
        7.0,
        4.4,
        [
            "Zero-CAC distribution into FNB SME.",
            "Unit-economics data banks never see.",
            "Prices a segment FNB cannot price today.",
            "Pipeline of credit-ready, thin-file customers.",
            "Better risk data on informal traders.",
        ],
        size=16,
        gap=18,
    )
    rect(s, 8.3, 1.95, 4.2, 4.35, fill=BRAND, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
    tf = textbox(s, 8.62, 2.25, 3.56, 3.8, anchor=MSO_ANCHOR.MIDDLE)
    para(tf, "THE FNB CASE", size=11.5, color=SKY, bold=True, first=True, align=PP_ALIGN.CENTER, space_after=14)
    para(
        tf,
        "Qualified, credit-ready small-business customers — with the risk data to match.",
        size=17,
        color=WHITE,
        bold=True,
        align=PP_ALIGN.CENTER,
        space_after=0,
        line_spacing=1.15,
    )

    # -- 13 · Closing -------------------------------------------------------
    s = hero_slide(prs)
    logo(s, 6.07, 1.5, 1.2, card_size=1.5)
    tf = textbox(s, 1.5, 3.25, 10.33, 1.2)
    para(tf, "Thank you", size=54, color=WHITE, bold=True, first=True, align=PP_ALIGN.CENTER, space_after=0)
    tf = textbox(s, 1.5, 4.4, 10.33, 0.5)
    para(
        tf,
        "Code calculates · AI explains · The paper trail is a by-product.",
        size=18,
        color=MIST,
        first=True,
        align=PP_ALIGN.CENTER,
        space_after=0,
    )
    chip(s, 4.67, 5.2, 4.0, 0.55, "grow-box.vercel.app", fill=MIST, color=BRAND, size=16)

    prs.save(out_path)
    return prs


# ------------------------------------------------------------- verification ---
def audit(prs: Presentation) -> list[str]:
    """Rough overflow check: flag text that likely does not fit its box."""
    warnings: list[str] = []
    for index, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            tf = shape.text_frame
            if not tf.text.strip():
                continue
            total_pt = 0.0
            for paragraph in tf.paragraphs:
                sizes = [run.font.size.pt for run in paragraph.runs if run.font.size]
                size = max(sizes) if sizes else 14.0
                text = "".join(run.text for run in paragraph.runs)
                width_pt = (shape.width / 914400 - 0.15) * 72
                chars_per_line = max(6, int(width_pt / (0.55 * size)))
                lines = max(1, -(-len(text) // chars_per_line))
                total_pt += lines * size * 1.25
            margin_top = tf.margin_top if tf.margin_top is not None else 45720
            margin_bottom = tf.margin_bottom if tf.margin_bottom is not None else 45720
            capacity_pt = shape.height / 12700 - (margin_top + margin_bottom) / 12700
            if total_pt > capacity_pt:
                snippet = tf.text.strip().replace("\n", " ")[:48]
                warnings.append(
                    f"  slide {index}: needs ~{total_pt:.0f}pt, fits {capacity_pt:.0f}pt — '{snippet}'"
                )
    return warnings


def main() -> int:
    out_path = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_OUT
    if not LOGO.exists():
        print(f"ERROR: logo not found at {LOGO}")
        return 1
    prs = build(out_path)
    reopened = Presentation(out_path)
    slide_count = len(reopened.slides)
    titles = []
    for slide in reopened.slides:
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text.strip():
                titles.append(shape.text.strip().splitlines()[0])
                break
    warnings = audit(prs)
    print(f"Wrote {out_path} — {slide_count} slides")
    for number, title in enumerate(titles, start=1):
        print(f"  {number:2d}. {title}")
    if warnings:
        print(f"Overflow warnings ({len(warnings)}):")
        for warning in warnings:
            print(warning)
    else:
        print("Overflow check: clean")
    return 0


if __name__ == "__main__":
    sys.exit(main())
